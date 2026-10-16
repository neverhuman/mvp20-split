"""Replot existing v1 evidence; never train, execute or modify the frozen baseline."""
from pathlib import Path
import json
import hashlib
import base64
import shutil
import fitz
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

HERE = Path(__file__).resolve().parent
BASE = Path('Z:/mvp20-split/v1')
ORIGINAL = Path('Z:/mvp20-split/v2/docs/Images/V1_Evidence')
INK, MUTED, ACCENT, PALE, LINE = '15181C', '606672', 'EF462D', 'FBE7E2', 'D9DADD'
GRAY = 'A5AAB2'
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(16), Inches(9)
pdf = fitz.open()
provenance = []


def load(path):
    raw = (BASE / path).read_bytes()
    provenance.append({'path': 'v1/' + path, 'sha256': hashlib.sha256(raw).hexdigest()})
    d = json.loads(raw)
    if 'signature' in d and isinstance(d['signature'], dict) and 'payload' in d['signature']:
        return json.loads(base64.b64decode(d['signature']['payload']))
    return d


density_path = 'amphora-molecule-property/reference/mvp20/data/radonpy_pi1070_density_ml_and_van_krevelen_physics_receipt.v1.json'
den = load(density_path)['kpis']
c8_path = 'amphora-molecule-property/benchmarks/phys-ml-stack-comparison-c8-v3-final/receipt_body.canonical.json'
c8 = load(c8_path)
tg = c8['per_property_comparison']['tg']
c7_path = 'amphora-molecule-property/benchmarks/tg-density-external-validity-c7-v3-final/receipt_body.canonical.json'
c7 = load(c7_path)
sol_path = 'amphora-solvency-score/evidence/audit/hansen-hardened-validation-v1/bigsoldb-admissibility-hardened-lso-receipt.v1.dsse.json'
sol = load(sol_path)
stack_path = 'amphora-molecule-property/reference/mvp20/data/stacked_ml_and_van_krevelen_physics_density_radonpy_receipt.v1.json'
stack = load(stack_path)['kpis']


def color(s):
    return tuple(int(s[i:i+2], 16) / 255 for i in (0, 2, 4))


class Slide:
    def __init__(self, number, kicker, title, subtitle):
        self.slide = prs.slides.add_slide(prs.slide_layouts[6])
        self.page = pdf.new_page(width=1152, height=648)
        self.rect(.43, .48, .09, .28, ACCENT)
        self.text(.66, .45, 14.7, .3, kicker.upper(), 11, ACCENT, True)
        self.text(.43, 1.05, 15.1, 1.15, title, 32, INK, True)
        self.text(.43, 2.25, 14.9, .65, subtitle, 16, MUTED)
        self.text(.43, 8.56, 14.7, .25, 'Amphora Labs  /  Confidential  /  v1 evidence  /  October 2026', 10, MUTED)
        self.text(15.05, 8.56, .5, .25, str(number), 10, MUTED)

    def rect(self, x, y, w, h, fill):
        shape = self.slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        shape.fill.solid(); shape.fill.fore_color.rgb = RGBColor.from_string(fill)
        shape.line.fill.background()
        self.page.draw_rect(fitz.Rect(x*72, y*72, (x+w)*72, (y+h)*72), color=None, fill=color(fill))

    def text(self, x, y, w, h, value, size=16, fill=INK, bold=False):
        box = self.slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        for n, line in enumerate(value.split('\n')):
            p = tf.paragraphs[0] if n == 0 else tf.add_paragraph()
            p.text = line
            p.font.name = 'Arial'; p.font.size = Pt(size); p.font.bold = bold
            p.font.color.rgb = RGBColor.from_string(fill)
            p.space_after = Pt(0)
        remaining = self.page.insert_textbox(fitz.Rect(x*72, y*72, (x+w)*72, (y+h)*72), value,
            fontsize=size, fontname='hebo' if bold else 'helv', color=color(fill), lineheight=1.13)
        if remaining < 0:
            raise ValueError(f'Text overflow ({remaining}): {value}')

    def source(self, text):
        self.text(.43, 7.83, 15, .58, text, 10, MUTED)

    def panel(self, headline, body, conclusion):
        self.rect(11.25, 3.25, 4.3, 4.18, PALE)
        self.text(11.55, 3.54, 3.7, .85, headline, 24, ACCENT, True)
        self.text(11.55, 4.65, 3.65, 1.5, body, 17)
        self.text(11.55, 6.35, 3.65, .75, conclusion, 14, MUTED)

    def table(self, headers, rows, widths):
        x0 = .43; y = 3.31
        for txt, w in zip(headers, widths):
            self.text(x0, y, w-.13, .65, txt, 13, MUTED, True); x0 += w
        self.rect(.43, 4.02, sum(widths), .014, INK)
        y = 4.10
        for i, row in enumerate(rows):
            highlighted = i == len(rows)-1
            if highlighted: self.rect(.43, y-.02, sum(widths), 1.05, PALE)
            x0 = .43
            for txt, w in zip(row, widths):
                self.text(x0+.06, y+.13, w-.17, .86, txt, 16, ACCENT if highlighted else INK, highlighted)
                x0 += w
            self.rect(.43, y+1.04, sum(widths), .01, LINE)
            y += 1.08


s = Slide(1, 'Polymer density / same-test-set comparison',
    'v1 learned polymer density beyond\nclassical group-contribution estimates',
    '1,077 polymer structures from RadonPy PI1070; the model and physics baseline share the same 216-polymer test set.')
b = den['physics_baseline']['vk_test_slice_MAE_g_per_cm3']; m = den['ml_baseline']['test_MAE_g_per_cm3']
lift = (1-m/b)*100
s.table(['Approach', 'How it was tested', 'MAE (g/cm3)', 'Fit (R2)'], [
    ['Van Krevelen physics', 'Random 60/20/20\ntrain / calibrate / test', f'{b:.4f}', '0.026'],
    ['Mean baseline', 'Same test polymers', '0.1268', '-'],
    ['Amphora v1 / random forest', '646 train / 215 calibrate\n216 test', f'{m:.4f}', '0.720']], [3.15,3.52,1.8,1.75])
s.panel(f'{lift:.1f}% less error', 'Learned molecular features improve density estimates over the named physics baseline.',
    'Supports screening within the tested chemical distribution.')
s.source('Source: RadonPy density comparison receipt [D1]. MD-simulated density at 300 K; random split.\nWithin-corpus result; novel-chemistry transfer is not established [A1].')

s = Slide(2, 'Thermal behavior / experimental polymer data',
    'v1 benchmarked thermal predictions\nagainst experimental polymer measurements',
    'OpenPoly glass-transition temperature: a matched 89-polymer test set compares physics, learned chemistry and their combination.')
tgp = tg['physics']['test_MAE_K']; tgm = tg['ml']['test_MAE_K']
s.table(['Approach', 'How it was tested', 'MAE (K)', 'Fit (R2)'], [
    ['Transferred physics model', 'Van Krevelen Ridge;\nno OpenPoly refitting', f'{tgp:.2f}', '-0.245'],
    ['Physics + ML stack', 'Same train / calibrate /\ntest partition', '54.05', '0.363'],
    ['Amphora v1 / random forest', '265 train / 88 calibrate\n89 test', f'{tgm:.2f}', '0.405']], [3.15,3.52,1.8,1.75])
s.panel(f'{(1-tgm/tgp)*100:.1f}% less error', 'On this matched test set, learned chemistry outperformed the transferred physics model.',
    'A measured starting point for prioritizing thermal-property experiments.')
s.source('Source: final physics / ML / stack comparison receipt [T1]. Experimental OpenPoly Tg; fresh random 60/20/20 split.\nBaseline was fitted on a different, simulated corpus. Single-corpus benchmark; not prospective lab validation [A1].')

s = Slide(3, 'Formulation support / unseen-solute testing',
    'v1 improved solubility predictions\nfor molecules withheld from training',
    'BigSolDB: 94,532 measurement groups across 1,372 solutes, evaluated with leave-solute-out cross-validation.')
full = sol['full_corpus_signed_comparator']; hardened = sol['admissibility_hardened_evaluation']
s.table(['Approach', 'How it was tested', 'RMSE (log S)'], [
    ['Best naive baseline', 'Solvent-mean baseline\non the same evaluation', f"{full['best_baseline_rmse']:.3f}"],
    ['Data-cleaning stress test', 'Contradictory state-point\ngroups removed; same protocol', f"{hardened['model_rmse']:.3f}"],
    ['Amphora v1', 'Test solutes withheld\nfrom training', f"{full['model_rmse']:.3f}"]], [3.15,4.92,2.15])
s.panel('25.0% less error', 'The advantage held after removing contradictory measurements: RMSE 0.877 became 0.874.',
    'Relevant to solvent selection; this is not a drug-release benchmark.')
s.source('Source: BigSolDB admissibility-hardened leave-solute-out audit [S1]. Full-corpus RMSE 95% CI: 0.842-0.915.\nLiterature-data benchmark and internal robustness audit; no prospective formulation or release claim.')

s = Slide(4, 'Evidence process / repeatable technical diligence',
    'v1 built an evidence trail\nbehind the performance numbers',
    'Named baselines, retained test protocols and explicit failure findings make the results inspectable.')
cards = [
    ('01', 'Compare fairly', 'Density and thermal comparisons use matched test sets. Solubility holds entire solutes out of training.'),
    ('02', 'Challenge the data', 'Removing contradictory solubility measurements preserved the gain: RMSE 0.877 to 0.874.'),
    ('03', 'Keep the failures', 'Hybrid density improved RMSE, but not MAE. Novel-chemistry transfer failed; its limits remain explicit.')]
for i,(n,title,body) in enumerate(cards):
    x=.43+i*5.13
    s.rect(x,3.38,4.86,3.44,'F5F5F6')
    s.text(x+.22,3.61,4.4,.5,n,25,ACCENT,True)
    s.text(x+.22,4.28,4.4,.5,title,23,INK,True)
    s.text(x+.22,5.02,4.37,1.48,body,17)
s.text(.43,7.10,15,.5,'The investment case: an existing, inspectable foundation for choosing the next experiments.',21,INK,True)
s.source('Sources: comparison receipts [D1, T1], solubility stress test [S1], density stack [D2], external-validity audit [A1].\nRetrospective computational evidence. Source receipts retain their original audit status; signatures were not re-verified for this deck.')

notes = '''V1 EVIDENCE SLIDES — SOURCE AND SPEAKER NOTES
Prepared 2026-10-16. Presentation styling follows ComparisonSlide.png (white, charcoal, coral highlight).
All charts/tables replot retained v1 results. No new training, experiment, signature verification or lab validation was performed.

SLIDE 1 / DENSITY
MAE 0.1327425203 -> 0.0566447055 g/cm3, a calculated 57.327% reduction.
RadonPy labels are molecular-dynamics simulations, not measured production samples.
Random within-corpus split: 646 train / 215 calibration / 216 test. R2 is not Pearson r.
The source dataset name PI1070 does not determine the retained 1,077-row evaluation count.

SLIDE 2 / THERMAL
Uses the fresh, internally matched C8 comparison: 80.2999166 -> 48.7735715 K MAE, a calculated 39.261% reduction.
265 train / 88 calibration / 89 test = 442 valid experimental OpenPoly records.
The physics baseline was transferred from a simulated Polyuniverse fit without OpenPoly refitting.
This is not a universal win over all physics models or state of the art.
Do not mix this row with the separate 51.2375 K scaffold-LOSO result or its different test protocol.
The final external-validity audit refutes novel-chemistry transfer; these slides preserve a within-corpus claim only.
The Polyuniverse absolute-temperature unit issue is not used as a positive performance claim here.

SLIDE 3 / SOLUBILITY
Full-corpus model RMSE 0.8770255643 vs solvent-mean 1.1686915142 (24.96% reduction).
95% cluster-bootstrap interval 0.8416341721–0.9154344725.
Hardened subset: 94,015 groups / 1,372 solutes; RMSE 0.8741657377 vs 1.1667252776.
517 contradictory state-point groups (1,193 raw rows) excluded from both train/test material.
This is small-molecule solubility supporting formulation research, not polymer dissolution or drug-release validation.

SLIDE 4 / EVIDENCE PROCESS
Density stack same-run RMSE 0.0723218955 vs RF 0.0861986512, but stack MAE 0.0569032315 vs RF 0.0566934387.
This demonstrates why choosing only a flattering metric is insufficient. No blanket claim that physics+ML always wins.
Audit-only development signatures are not external review or scientific validation. No signing-trust claim is made.
Dielectric positive figures from the old pack were omitted because C8 retracts the contaminated-population claim.
Creative license is limited to titles, layout and investment framing. No invented data, scatter points, confidence intervals or competitor results.

SOURCE MAP (repository-relative; SHA256 pins are in source_manifest.json)
'''
for label,path in [('D1',density_path),('D2',stack_path),('T1',c8_path),('S1',sol_path),('A1',c7_path)]:
    notes += f'{label}: v1/{path}\n'
for slide in prs.slides:
    slide.notes_slide.notes_text_frame.text = notes
prs.save(HERE/'V1_Polymer_Evidence.pptx')
pdf.set_metadata({'title':'Amphora v1 — Polymer Evidence','author':'Amphora Labs','subject':'Retrospective benchmark evidence'})
pdf.save(HERE/'V1_Polymer_Evidence.pdf')
for i,page in enumerate(pdf):
    page.get_pixmap(matrix=fitz.Matrix(1.5,1.5),alpha=False).save(HERE/f'Slide_{i+1:02}.png')
(HERE/'Source_Notes.txt').write_text(notes,encoding='utf-8')
(HERE/'source_manifest.json').write_text(json.dumps(provenance,indent=2)+'\n',encoding='utf-8')
ORIGINAL.mkdir(parents=True,exist_ok=True)
if HERE != ORIGINAL:
    for p in HERE.iterdir():
        if p.is_file(): shutil.copy2(p,ORIGINAL/p.name)
print(f'Created {len(prs.slides)} editable slides, PDF, PNGs and source notes in {HERE}; copied to {ORIGINAL}')
