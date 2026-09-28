"""Three visual investor slides, condensed from the full positioning deck."""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN
import math

ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'docs/Artifacts/AMPHORA_Engine_VC_Visual_v3.pptx'
prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
BG='101D2C'; PANEL='192B3E'; PANEL2='22394D'; WHITE='F1F5F6'; MUTED='AFC1CE'
TEAL='56D9BD'; BLUE='83BFFF'; GOLD='F4C77E'; PURPLE='BEA5F4'; LINE='40586A'
SOURCES={
'S1':('Benchling AI','https://www.benchling.com/ai'),
'S2':('Citrine Catalyst','https://citrine.io/platform/citrine-catalyst/'),
'S3':('Citrine sequential learning','https://citrine.io/doe/'),
'S4':('Schrödinger active learning','https://www.schrodinger.com/platform/products/active-learning-applications/'),
'S5':('Synthace','https://www.synthace.com/what-we-do'),
'S6':('Automata LINQ','https://docs.automata.tech/tasks-transport'),
'S7':('Recursion OS','https://www.recursion.com/faq'),
'S8':('Claude Science','https://www.anthropic.com/news/claude-science-ai-workbench')}

def rgb(c): return RGBColor.from_string(c)
def box(s,x,y,w,h,fill=PANEL,stroke=None):
 sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h));sh.adjustments[0]=.12
 sh.fill.solid();sh.fill.fore_color.rgb=rgb(fill)
 if stroke: sh.line.color.rgb=rgb(stroke)
 else: sh.line.fill.background()
 return sh
def text(s,x,y,w,h,t,size=18,color=WHITE,bold=False,align=PP_ALIGN.LEFT,url=None):
 sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=sh.text_frame;tf.clear();tf.word_wrap=True
 tf.margin_top=tf.margin_bottom=tf.margin_left=tf.margin_right=0
 for i,row in enumerate(t.split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=row;p.font.name='Aptos';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=rgb(color);p.alignment=align;p.space_after=Pt(5)
  if url:
   for r in p.runs:r.hyperlink.address=url
 return sh
def line(s,x,y,xx,yy,c=LINE,width=1.5):
 sh=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x),Inches(y),Inches(xx),Inches(yy));sh.line.color.rgb=rgb(c);sh.line.width=Pt(width)
def arrow(s,x,y,xx,yy,c=TEAL):
 line(s,x,y,xx,yy,c,2);a=math.atan2(yy-y,xx-x)
 for d in [-.52,.52]:line(s,xx,yy,xx-.11*math.cos(a+d),yy-.11*math.sin(a+d),c,2)
def dot(s,x,y,c=TEAL,d=.12):
 sh=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x-d/2),Inches(y-d/2),Inches(d),Inches(d));sh.fill.solid();sh.fill.fore_color.rgb=rgb(c);sh.line.fill.background()
def chip(s,x,y,w,t,c=TEAL):
 box(s,x,y,w,.31,PANEL2);text(s,x+.10,y+.06,w-.2,.22,t,10,c,True)
def base(kicker,title,sub='',state='INVESTMENT THESIS'):
 s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=rgb(BG)
 text(s,.5,.30,12,.25,kicker.upper(),10,TEAL,True);text(s,.5,.80,12.3,.90,title,30,WHITE,True)
 if sub:text(s,.52,1.83,12.1,.56,sub,16,MUTED)
 line(s,.5,7.03,12.83,7.03,LINE,.65)
 text(s,.5,7.16,7,.19,'AMPHORA  /  CONFIDENTIAL INVESTOR DISCUSSION  /  SEPTEMBER 2026',8,MUTED)
 text(s,9.15,7.14,3.05,.22,state,8,TEAL,align=PP_ALIGN.RIGHT)
 text(s,12.45,7.10,.35,.29,f'{len(prs.slides):02}',10,WHITE,True,PP_ALIGN.RIGHT)
 return s
def band(s,t,c=TEAL,y=6.36):
 box(s,.52,y,12.29,.43,PANEL2);text(s,.72,y+.085,11.9,.31,t,14,c,True)
def notes(s,t,ids=()):
 body=t+'\n\nAmphora architecture is a proposed product direction. Existing numerical search/evidence code has been source-reviewed; full graph execution, biological evaluators and pathway evolution are not established as delivered. No commercial performance or exclusive novelty is claimed.\nLocal basis: docs/Architecture/Requirements/Graph_Native_Research_Engine.md; Product_Requirements.md; Evolutionary_Code_Reuse_Baseline.md.\nCompetitive diligence: docs/Presentations/VC_Positioning_Research.md.\n'
 for key in ids:body+=f'\n[{key}] {SOURCES[key][0]}: {SOURCES[key][1]} (accessed 2026-09-16; vendor description)'
 s.notes_slide.notes_text_frame.text=body
def refs(s,ids,y=6.88):
 x=.54
 for key in ids:
  label=f'{key} {SOURCES[key][0]}'
  w=max(1.25,len(label)*.060)
  text(s,x,y,w,.14,label,7.5,MUTED,url=SOURCES[key][1]);x+=w+.17

def step(s,x,y,w,title,sub='',c=TEAL,h=.95):
 box(s,x,y,w,h,PANEL,c)
 text(s,x+.13,y+.18,w-.26,.35,title,17,WHITE,True,PP_ALIGN.CENTER)
 if sub:text(s,x+.12,y+.61,w-.24,.32,sub,11,MUTED,align=PP_ALIGN.CENTER)

def both(s,x,y,xx,yy,c=TEAL):
 arrow(s,x,y,xx,yy,c);arrow(s,xx,yy,x,y,c)

def graph(s,x,y,w,h,points,edges,highlight=(),radius=.10):
 for a,b in edges:
  ax,ay=points[a];bx,by=points[b]
  c=TEAL if (a,b) in highlight else LINE
  line(s,x+ax*w,y+ay*h,x+bx*w,y+by*h,c,2.8 if c==TEAL else 1.1)
 for i,(px,py) in enumerate(points):
  c=TEAL if any(i in e for e in highlight) else BLUE
  dot(s,x+px*w,y+py*h,c,radius)

def person(s,x,y,c=BLUE):
 dot(s,x,y,c,.13);line(s,x,y+.10,x,y+.32,c,2.5)
 line(s,x-.12,y+.18,x+.12,y+.18,c,2)
 line(s,x,y+.32,x-.10,y+.44,c,2);line(s,x,y+.32,x+.10,y+.44,c,2)

def mini(s,x,y,w,h,variant=0,c=TEAL):
 pts=[(.06,.50),(.36,.18),(.36,.82),(.67,.50),(.95,.50)]
 edges=[(0,1),(0,2),(1,3),(2,3),(3,4)]
 if variant:pts.append((.68,.02));edges=[(0,1),(0,2),(1,5),(5,4),(2,3),(3,4)]
 for a,b in edges:arrow(s,x+pts[a][0]*w,y+pts[a][1]*h,x+pts[b][0]*w,y+pts[b][1]*h,c)
 for px,py in pts:box(s,x+px*w-.12,y+py*h-.12,.24,.24,PANEL2,c)

# Slide 1: context becomes an actionable research map.
s=base('01 / Why build','Make the whole problem actionable.',state='PRODUCT VISION')
text(s,.70,2.08,4.9,.40,'GROWING CONTEXT',17,MUTED,True)
text(s,7.10,2.08,5.0,.40,'LIVING RESEARCH MAP',17,TEAL,True)
pts=[(.02,.25),(.16,.07),(.19,.43),(.04,.74),(.31,.26),(.36,.62),(.24,.94),(.50,.06),(.55,.36),(.48,.83),(.71,.20),(.72,.68),(.94,.08),(.94,.46),(.88,.94)]
edges=[(0,1),(0,2),(0,3),(1,4),(1,7),(2,4),(2,5),(3,5),(3,6),(4,7),(4,8),(5,8),(5,9),(6,9),(7,10),(8,10),(8,11),(9,11),(10,12),(10,13),(11,13),(11,14),(13,14)]
graph(s,.82,2.92,4.32,2.40,pts,edges,radius=.10)
for x,y,t,c in [(.69,2.64,'Data',BLUE),(3.54,2.64,'Models',PURPLE),(.73,5.55,'Experiments',TEAL),(3.64,5.55,'Expertise',GOLD)]:text(s,x,y,1.7,.34,t,16,c,True)
arrow(s,5.38,4.12,6.54,4.12,TEAL)
text(s,5.27,4.44,1.46,.27,'AMPHORA',11,TEAL,True,PP_ALIGN.CENTER)
box(s,6.90,2.72,5.68,3.12,PANEL,LINE)
graph(s,7.18,3.01,5.02,2.40,pts,edges,[(0,2),(2,5),(5,8),(8,10),(10,13)],.10)
for x,y,t,c in [(7.23,3.34,'Ask',BLUE),(8.85,4.61,'Test',TEAL),(10.63,3.03,'Learn',GOLD)]:
 box(s,x,y,1.04,.51,PANEL2,c);text(s,x+.1,y+.12,.84,.30,t,16,c,True,PP_ALIGN.CENTER)
text(s,.74,6.35,11.88,.49,'See the connections. Choose the next step.',27,WHITE,True,PP_ALIGN.CENTER)
notes(s,'WHY BUILD — talk track, about 45 seconds. Point to the left: research creates papers, datasets, measurements, custom models and expert assumptions. The problem is maintaining the relationships and turning that context into decisions as it changes. Point to the right: Amphora is designed to make the context navigable and executable; the highlighted path is the work we choose to do next, while the surrounding evidence stays available. This is the reason for a shared engine: less repeated reconstruction, more continuity from question to experiment to learning. The two views are a conceptual transformation, not an assertion that every competitor lacks connected data. The performance benefit remains to be demonstrated. Massive context lives in durable graph/artifact storage; relevant subsets are projected into LLM context.')

# Slide 2: the full architectural relationship, mostly drawn rather than described.
s=base('02 / How it works','One visual workspace. An adaptive research engine.',state='TARGET ARCHITECTURE')
for i,(t,c) in enumerate([('Scientist',BLUE),('Model builder',PURPLE),('Reviewer',GOLD),('Program lead',TEAL)]):
 x=3.17+i*2.42;person(s,x,1.97,c);text(s,x+.25,1.94,1.91,.32,t,14,MUTED)
box(s,.53,2.60,12.28,3.80,PANEL,LINE)
text(s,.76,2.79,11.80,.29,'VISUALIZE  ·  INSPECT  ·  EDIT  ·  RUN',12,TEAL,True)
for i,(t,c) in enumerate([('Data sources',BLUE),('Custom models',PURPLE),('Genetic algorithms',TEAL)]):
 y=3.45+i*.73;box(s,.79,y,2.36,.55,PANEL2,c);text(s,.92,y+.14,2.10,.32,t,14,c,True,PP_ALIGN.CENTER)
 line(s,3.19,y+.28,3.46,y+.28,c)
line(s,3.46,3.73,3.46,5.19,LINE);arrow(s,3.46,4.48,3.92,4.48)
text(s,4.06,3.39,2.64,.38,'Execution graph',18,TEAL,True,PP_ALIGN.CENTER)
mini(s,4.16,4.12,2.37,1.06)
both(s,6.66,4.48,7.12,4.48,TEAL)
dot(s,7.91,4.48,BLUE,1.31)
text(s,7.30,4.23,1.22,.48,'LLM',28,BG,True,PP_ALIGN.CENTER)
both(s,8.64,4.48,9.10,4.48,PURPLE)
text(s,9.24,3.29,3.19,.70,'Neuro-symbolic\nknowledge graph',18,PURPLE,True,PP_ALIGN.CENTER)
kp=[(.07,.55),(.28,.13),(.40,.72),(.60,.35),(.73,.90),(.96,.58),(.91,.08)]
graph(s,9.52,4.12,2.58,1.33,kp,[(0,1),(0,2),(1,3),(2,3),(2,4),(3,5),(3,6),(4,5),(5,6)],[(0,2),(2,3)],.13)
text(s,4.05,5.90,8.20,.28,'Evidence + execution history',13,GOLD,True,PP_ALIGN.CENTER)
text(s,.77,6.62,11.84,.28,'Materials  ↔  Chemistry  ↔  Biology',16,MUTED,align=PP_ALIGN.CENTER)
notes(s,'HOW — talk track, about 60 seconds. Start with the three inputs: data sources supply evidence, custom models supply scientific evaluations, and genetic algorithms explore candidate designs. The execution graph composes these into steps, branches and repeat loops. The LLM works between that graph and the neuro-symbolic knowledge graph: it interprets objectives, retrieves relevant context, explains outcomes and proposes changes. The knowledge graph carries mechanisms, evidence, symbolic rules, uncertainty and applicability. All of this is inside one visual workspace, where scientists, model builders, reviewers and leads navigate views suited to their roles and customize campaigns. The LLM does not receive unlimited context or directly activate unreviewed amendments; it uses bounded projections and reviewed changes. The displayed arrows describe architecture, not an implemented product. Execution and provenance history is derived. Genetic parameter search exists in the code; general graph-topology evolution remains future work.')

# Slide 3: three visual reasons the combined design matters.
s=base('03 / The distinctive combination','Human control. Compounding intelligence.',state='DESIGN GOALS / TO PROVE')
for x in [4.44,8.76]:line(s,x,2.36,x,6.16,LINE,.8)
text(s,.68,2.12,3.57,.42,'NAVIGATE CONTEXT',16,PURPLE,True,PP_ALIGN.CENTER)
text(s,4.71,2.12,3.80,.42,'SHAPE EXECUTION',16,BLUE,True,PP_ALIGN.CENTER)
text(s,9.03,2.12,3.62,.42,'EVOLVE PATHWAYS',16,TEAL,True,PP_ALIGN.CENTER)
# A selected neighborhood and a magnified, editable subgraph.
graph(s,.86,2.99,3.10,1.62,pts,edges,radius=.07)
sel=box(s,1.94,3.44,1.03,.68,BG,PURPLE);sel.fill.background()
line(s,1.94,4.12,1.53,4.91,PURPLE);line(s,2.97,4.12,3.41,4.91,PURPLE)
box(s,1.34,4.93,2.23,.93,PANEL,PURPLE)
graph(s,1.58,5.09,1.70,.52,[(0,.5),(.45,.05),(.5,.95),(1,.5)],[(0,1),(0,2),(1,3),(2,3)],[(0,1),(1,3)],.12)
# A graph edit with a new branch and human action.
mini(s,4.97,3.10,3.18,1.54,0,BLUE)
box(s,6.10,5.09,1.46,.52,PANEL2,GOLD);text(s,6.20,5.22,1.26,.30,'New assay',14,GOLD,True,PP_ALIGN.CENTER)
line(s,6.12,4.48,6.12,4.82,GOLD);arrow(s,6.12,4.82,6.67,5.04,GOLD)
line(s,7.08,5.08,7.61,4.18,GOLD)
# Editable cursor icon.
cursor=s.shapes.add_shape(MSO_SHAPE.RIGHT_TRIANGLE,Inches(7.79),Inches(5.25),Inches(.25),Inches(.39));cursor.rotation=330;cursor.fill.solid();cursor.fill.fore_color.rgb=rgb(WHITE);cursor.line.fill.background()
# A parent pathway and an evolved descendant.
mini(s,9.32,2.92,2.84,.87,0,LINE)
arrow(s,10.85,4.01,10.85,4.57,TEAL)
mini(s,9.32,4.84,2.84,1.03,1,TEAL)
for x,t,c in [(.64,'Massive context',PURPLE),(4.78,'Role-based editing',BLUE),(9.02,'Search the research',TEAL)]:text(s,x,6.11,3.63,.39,t,20,c,True,PP_ALIGN.CENTER)
text(s,.72,6.71,11.90,.27,'Reusable science. Customizable campaigns. Retained learning.',17,WHITE,True,PP_ALIGN.CENTER)
notes(s,'WHY THE COMBINATION MATTERS — talk track, about 60 seconds. Left: users zoom from the whole research context into the exact evidence or mechanism needed for a decision. Center: authorized users change a campaign directly, including new data/model/assay steps, without making every pathway a new orchestration implementation. Different roles have different views and permissions over shared state. Right: genetic and LLM-assisted processes eventually explore campaign structure as well as candidate parameters, retaining ancestry and the outcomes of each attempt. Existing algorithms, workflow graphs, knowledge graphs and visual editors have prior art; differentiation is the scientific integration and how effectively it supports real research. Custom modules may still require coding and validation. The test is measured decision quality, adaptation effort and reuse, not novelty by assertion. Broader workflow evolution has prior art in AFlow https://arxiv.org/abs/2410.10762 and ADAS https://arxiv.org/abs/2408.08435. No present implementation or superior commercial performance is implied.',tuple(SOURCES))

prs.core_properties.title='Amphora | Visual Engine Story'
prs.core_properties.subject='Presenter-led visual story: why, how, and differentiation'
prs.core_properties.author='Amphora'
prs.save(OUT)
print(f'Saved {len(prs.slides)} slides: {OUT}')
