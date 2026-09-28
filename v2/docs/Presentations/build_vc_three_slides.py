"""Three visual investor slides, condensed from the full positioning deck."""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN
import math

ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'docs/Artifacts/AMPHORA_Engine_VC_3_Slides_v2.pptx'
prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
BG='101D2C'; PANEL='192B3E'; PANEL2='22394D'; WHITE='F1F5F6'; MUTED='AFC1CE'
TEAL='56D9BD'; BLUE='83BFFF'; GOLD='F4C77E'; PURPLE='BEA5F4'; LINE='40586A'
SOURCES={
'S1':('Benchling AI','https://www.benchling.com/ai'),
'S2':('Citrine Catalyst','https://citrine.io/platform/citrine-catalyst/'),
'S3':('Citrine sequential learning','https://citrine.io/doe/'),
'S4':('Schrödinger active learning','https://www.schrodinger.com/platform/products/active-learning-applications/'),
'S5':('Synthace','https://www.synthace.com/what-we-do'),
'S6':('Automata LINQ','https://docs.automata.tech/tasks-transport'),
'S7':('Recursion OS','https://www.recursion.com/faq'),
'S8':('Claude Science','https://www.anthropic.com/news/claude-science-ai-workbench')}

def rgb(c): return RGBColor.from_string(c)
def box(s,x,y,w,h,fill=PANEL,stroke=None):
 sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h));sh.adjustments[0]=.12
 sh.fill.solid();sh.fill.fore_color.rgb=rgb(fill)
 if stroke: sh.line.color.rgb=rgb(stroke)
 else: sh.line.fill.background()
 return sh
def text(s,x,y,w,h,t,size=18,color=WHITE,bold=False,align=PP_ALIGN.LEFT,url=None):
 sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=sh.text_frame;tf.clear();tf.word_wrap=True
 tf.margin_top=tf.margin_bottom=tf.margin_left=tf.margin_right=0
 for i,row in enumerate(t.split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=row;p.font.name='Aptos';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=rgb(color);p.alignment=align;p.space_after=Pt(5)
  if url:
   for r in p.runs:r.hyperlink.address=url
 return sh
def line(s,x,y,xx,yy,c=LINE,width=1.5):
 sh=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x),Inches(y),Inches(xx),Inches(yy));sh.line.color.rgb=rgb(c);sh.line.width=Pt(width)
def arrow(s,x,y,xx,yy,c=TEAL):
 line(s,x,y,xx,yy,c,2);a=math.atan2(yy-y,xx-x)
 for d in [-.52,.52]:line(s,xx,yy,xx-.11*math.cos(a+d),yy-.11*math.sin(a+d),c,2)
def dot(s,x,y,c=TEAL,d=.12):
 sh=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x-d/2),Inches(y-d/2),Inches(d),Inches(d));sh.fill.solid();sh.fill.fore_color.rgb=rgb(c);sh.line.fill.background()
def chip(s,x,y,w,t,c=TEAL):
 box(s,x,y,w,.31,PANEL2);text(s,x+.10,y+.06,w-.2,.22,t,10,c,True)
def base(kicker,title,sub='',state='INVESTMENT THESIS'):
 s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=rgb(BG)
 text(s,.5,.30,12,.25,kicker.upper(),10,TEAL,True);text(s,.5,.80,12.3,.90,title,30,WHITE,True)
 if sub:text(s,.52,1.83,12.1,.56,sub,16,MUTED)
 line(s,.5,7.03,12.83,7.03,LINE,.65)
 text(s,.5,7.16,7,.19,'AMPHORA  /  CONFIDENTIAL INVESTOR DISCUSSION  /  SEPTEMBER 2026',8,MUTED)
 text(s,9.15,7.14,3.05,.22,state,8,TEAL,align=PP_ALIGN.RIGHT)
 text(s,12.45,7.10,.35,.29,f'{len(prs.slides):02}',10,WHITE,True,PP_ALIGN.RIGHT)
 return s
def band(s,t,c=TEAL,y=6.36):
 box(s,.52,y,12.29,.43,PANEL2);text(s,.72,y+.085,11.9,.31,t,14,c,True)
def notes(s,t,ids=()):
 body=t+'\n\nAmphora architecture is a proposed product direction. Existing numerical search/evidence code has been source-reviewed; full graph execution, biological evaluators and pathway evolution are not established as delivered. No commercial performance or exclusive novelty is claimed.\nLocal basis: docs/Architecture/Requirements/Graph_Native_Research_Engine.md; Product_Requirements.md; Evolutionary_Code_Reuse_Baseline.md.\nCompetitive diligence: docs/Presentations/VC_Positioning_Research.md.\n'
 for key in ids:body+=f'\n[{key}] {SOURCES[key][0]}: {SOURCES[key][1]} (accessed 2026-09-16; vendor description)'
 s.notes_slide.notes_text_frame.text=body
def refs(s,ids,y=6.88):
 x=.54
 for key in ids:
  label=f'{key} {SOURCES[key][0]}'
  w=max(1.25,len(label)*.060)
  text(s,x,y,w,.14,label,7.5,MUTED,url=SOURCES[key][1]);x+=w+.17

def step(s,x,y,w,title,sub='',c=TEAL,h=.95):
 box(s,x,y,w,h,PANEL,c)
 text(s,x+.13,y+.18,w-.26,.35,title,17,WHITE,True,PP_ALIGN.CENTER)
 if sub:text(s,x+.12,y+.61,w-.24,.32,sub,11,MUTED,align=PP_ALIGN.CENTER)

def both(s,x,y,xx,yy,c=TEAL):
 arrow(s,x,y,xx,yy,c);arrow(s,xx,yy,x,y,c)

# 1 — the reason to build, grounded in the problem the architecture solves.
s=base('01 / Why build Amphora','Research needs a system that holds the whole problem.','Connect large scientific context to a research plan that people can understand, change and execute.',state='PRODUCT VISION')
text(s,.62,2.57,3.05,.32,'THE CONTEXT KEEPS GROWING',12,MUTED,True)
labels=[('Papers + datasets',BLUE),('Experiments + failures',TEAL),('Models + predictions',PURPLE),('Constraints + expertise',GOLD)]
for i,(t,c) in enumerate(labels):
 y=3.11+i*.66;box(s,.58,y,3.05,.51,PANEL);text(s,.76,y+.12,2.68,.29,t,16,c,True)
 line(s,3.65,y+.25,4.15,4.40,c)
arrow(s,4.15,4.40,4.63,4.40)
box(s,4.68,2.99,3.77,3.18,PANEL2,TEAL)
text(s,4.93,3.26,3.27,.31,'AMPHORA',16,TEAL,True,PP_ALIGN.CENTER)
text(s,4.93,3.95,3.27,1.40,'A living research map\nthat can act',24,WHITE,True,PP_ALIGN.CENTER)
text(s,4.93,5.55,3.27,.32,'Context → decisions → execution',12,MUTED,align=PP_ALIGN.CENTER)
arrow(s,8.50,4.40,9.04,4.40)
for i,(h,b,c) in enumerate([('UNDERSTAND','See the relevant relationships',TEAL),('ADAPT','Change the plan as evidence changes',BLUE),('COMPOUND','Reuse what the campaign teaches us',GOLD)]):
 y=3.10+i*1.00;dot(s,9.27,y+.19,c,.12);text(s,9.51,y,3.20,.32,h,16,c,True);text(s,9.51,y+.44,3.19,.55,b,16,WHITE)
text(s,.62,6.61,12.09,.34,'WHY BUILD: make complex research manageable—and turn its context into executable decisions.',18,TEAL,True)
notes(s,'The reason to build is the combined context-and-control problem: research spans sources, experiments, models, constraints and human expertise. The goal is to make relationships navigable and translate selected context into executable research. Large context means durable, addressable and versioned information across graphs and artifacts, not loading an entire corpus into an LLM prompt. Benefits are design goals to validate, not measured superiority. The engine should reduce repeated context reconstruction and pathway reimplementation while retaining evidence and human control.')

# 2 — exact system relationship, with all components enclosed by the visual workspace.
s=base('02 / How the engine works','Compose the science. Reason over context. Execute the graph.','Genetic search, data and custom models become configurable scientific capabilities.',state='TARGET ENGINE ARCHITECTURE')
box(s,.53,2.53,12.28,4.12,PANEL,LINE)
text(s,.76,2.72,6.25,.29,'INTERACTIVE VISUALIZATION + COLLABORATION WORKSPACE',12,TEAL,True)
text(s,7.51,2.75,5.02,.24,'SCIENTISTS  ·  MODEL BUILDERS  ·  REVIEWERS  ·  LEADS',9,MUTED,True,PP_ALIGN.RIGHT)
for i,(t,c) in enumerate([('DATA SOURCES',BLUE),('CUSTOM MODELS',PURPLE),('GENETIC ALGORITHMS',TEAL)]):
 y=3.43+i*.75;box(s,.78,y,2.44,.56,PANEL2,c);text(s,.91,y+.17,2.18,.27,t,12,c,True,PP_ALIGN.CENTER)
 line(s,3.25,y+.28,3.51,y+.28,c)
line(s,3.51,3.71,3.51,5.21,LINE);arrow(s,3.51,4.48,3.89,4.48)
box(s,3.94,3.43,2.78,2.18,PANEL2,TEAL)
text(s,4.12,3.64,2.42,.85,'EXECUTION\nGRAPH',21,TEAL,True,PP_ALIGN.CENTER)
text(s,4.15,4.65,2.37,.65,'Compose · branch · iterate\nRetain state · run tools',12,WHITE,align=PP_ALIGN.CENTER)
both(s,6.77,4.48,7.26,4.48,TEAL)
box(s,7.31,3.74,1.44,1.52,PANEL2,BLUE);text(s,7.48,3.98,1.10,.45,'LLM',25,BLUE,True,PP_ALIGN.CENTER)
text(s,7.43,4.68,1.20,.35,'Reason + propose',10,WHITE,align=PP_ALIGN.CENTER)
both(s,8.80,4.48,9.29,4.48,PURPLE)
box(s,9.34,3.43,3.20,2.18,PANEL2,PURPLE)
text(s,9.50,3.65,2.88,.79,'NEURO-SYMBOLIC\nKNOWLEDGE GRAPH',17,PURPLE,True,PP_ALIGN.CENTER)
text(s,9.52,4.71,2.85,.64,'Mechanisms · evidence · rules\nUncertainty · applicability',12,WHITE,align=PP_ALIGN.CENTER)
text(s,3.93,5.89,8.49,.40,'Users inspect context, edit pathways and review proposed changes.',16,WHITE,True,PP_ALIGN.CENTER)
text(s,.77,6.35,11.78,.20,'SHARED FOUNDATION: versioned artifacts · source lineage · execution history · permissions · human gates',10,MUTED)
notes(s,'This is the requested system relationship: data sources, custom scientific models and genetic/evolutionary search are bounded capabilities composed by an executable campaign graph. The graph and LLM exchange current state, questions and proposed amendments; the LLM retrieves bounded context from and proposes changes to a separate neuro-symbolic domain graph. That graph represents entities, mechanisms, evidence, symbolic constraints, uncertainty and applicability; neural/model outputs remain evidence-scoped. The visual workspace surrounds all of these and provides role-appropriate views, graph editing, execution inspection, comparison and intervention. Massive context remains in versioned graphs/artifacts and is projected selectively into model context. LLM proposals do not directly rewrite trusted knowledge or activate execution: validation and authorized human gates apply. Genetic algorithms search candidate parameters; governed campaign-topology evolution is planned. Execution/provenance is derived history. This is a target architecture, not an implemented-system diagram.')

# 3 — concrete differentiating combination and why it matters.
s=base('03 / What makes the design distinctive','The differentiator is how the system learns—and stays usable.','Four capabilities designed to reinforce one another across materials, chemistry and biology.',state='DIFFERENTIATION THESIS')
cards=[(.57,2.63,'01','EVOLVE THE RESEARCH PATH','Search candidate designs and, over time,\nthe sequence of work used to discover them.',TEAL),
       (6.93,2.63,'02','REASON ACROSS MASSIVE CONTEXT','Navigate connected knowledge and evidence;\nbring the relevant context to each decision.',PURPLE),
       (.57,4.64,'03','CUSTOMIZE WITHOUT A FIXED PIPELINE','Mix data, genetic search and custom models\nby editing the executable campaign graph.',BLUE),
       (6.93,4.64,'04','LET EVERY ROLE SHAPE THE WORK','Scientists, developers and reviewers inspect,\ncustomize and intervene through visual views.',GOLD)]
for x,y,n,h,b,c in cards:
 box(s,x,y,5.84,1.79,PANEL)
 text(s,x+.19,y+.17,.51,.42,n,24,c,True)
 text(s,x+.88,y+.20,4.72,.34,h,12,c,True)
 text(s,x+.88,y+.69,4.70,.74,b,16,WHITE)
 # A small editable motif gives each differentiator a recognizable visual mark.
 if n=='01':
  for yy in [y+.79,y+1.24]:
   dot(s,x+.20,yy,c,.09);dot(s,x+.59,yy,c,.09);line(s,x+.25,yy,x+.54,yy,c)
  line(s,x+.59,y+.84,x+.20,y+1.19,c)
 elif n=='02':
  for dx,dy in [(.18,.88),(.63,.81),(.39,1.18),(.64,1.41)]:dot(s,x+dx,y+dy,c,.085)
  for a,b2 in [((.18,.88),(.39,1.18)),((.63,.81),(.39,1.18)),((.39,1.18),(.64,1.41))]:line(s,x+a[0],y+a[1],x+b2[0],y+b2[1],c)
 elif n=='03':
  for dy in [.80,1.05,1.30]:box(s,x+.16,y+dy,.50,.16,PANEL2,c)
 else:
  for dy in [.85,1.17]:dot(s,x+.24,y+dy,c,.11);line(s,x+.37,y+dy,x+.66,y+dy,c,2)
text(s,.63,6.67,12.04,.31,'THE VALUE TO PROVE: faster adaptation, better decisions and reusable scientific learning.',18,TEAL,True)
notes(s,'These are intended points of differentiation in combination, not claims that individual techniques are new or exclusive. 1: candidate optimization and future bounded topology evolution retain ancestry and reviewed constraints. 2: massive context is durable graph/artifact context with selective retrieval, not unlimited LLM memory; assertions preserve evidence, disagreement and applicability. 3: reviewed modules remain implementation code; users change compatible pathway composition through graphs, and new scientific capabilities can still require coding and validation. 4: Graph Studio supplies role-specific navigation, graph changes, approvals, history and interventions over shared state, not merely a static canvas. Existing platforms overlap with these features; prove advantage by comparing complete workflows. Evolutionary workflow design has prior art including AFlow https://arxiv.org/abs/2410.10762 and ADAS https://arxiv.org/abs/2408.08435. No broad novelty or patent claim is made. The near-term scientific/engineering evidence must establish utility across at least two distinct campaign representations.',tuple(SOURCES))

prs.core_properties.title='Amphora | Why Build the Research Engine'
prs.core_properties.subject='Three visual slides: why, how, and differentiation'
prs.core_properties.author='Amphora'
prs.save(OUT)
print(f'Saved {len(prs.slides)} slides: {OUT}')
