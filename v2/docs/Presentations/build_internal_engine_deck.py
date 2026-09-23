"""Build the editable internal engine explainer. Requires python-pptx and Pillow."""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'docs' / 'Artifacts' / 'AMPHORA_Research_Engine_Internal.pptx'
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BG='101D2C'; PANEL='192B3E'; PANEL2='22394D'; WHITE='F1F5F6'; MUTED='AFC1CE'
TEAL='56D9BD'; BLUE='83BFFF'; GOLD='F4C77E'; PURPLE='BEA5F4'; LINE='40586A'
FONT='Aptos'

def rgb(c): return RGBColor.from_string(c)
def box(s,x,y,w,h,fill=PANEL,line=None,r=0.12):
    sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if r else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    if r:
        sh.adjustments[0]=r
    sh.fill.solid(); sh.fill.fore_color.rgb=rgb(fill)
    sh.line.fill.background() if line is None else None
    if line: sh.line.color.rgb=rgb(line)
    return sh
def txt(s,x,y,w,h,text,size=18,color=WHITE,bold=False,align=PP_ALIGN.LEFT):
    sh=s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf=sh.text_frame; tf.clear(); tf.word_wrap=True
    tf.margin_left=tf.margin_right=0; tf.margin_top=tf.margin_bottom=0
    for i,line in enumerate(text.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.text=line; p.font.name=FONT; p.font.size=Pt(size); p.font.bold=bold; p.font.color.rgb=rgb(color)
        p.alignment=align; p.space_after=Pt(5)
    return sh
def line(s,x1,y1,x2,y2,color=LINE,width=1.8):
    sh=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    sh.line.color.rgb=rgb(color); sh.line.width=Pt(width)
    return sh
def arrow(s,x1,y1,x2,y2,color=TEAL):
    line(s,x1,y1,x2,y2,color,2)
    # Editable arrow tip, compatible with PowerPoint renderers.
    import math
    a=math.atan2(y2-y1,x2-x1)
    for d in [-0.52,0.52]:
        line(s,x2,y2,x2-.11*math.cos(a+d),y2-.11*math.sin(a+d),color,2)
def dot(s,x,y,d=.12,color=TEAL):
    sh=s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x-d/2), Inches(y-d/2), Inches(d), Inches(d))
    sh.fill.solid(); sh.fill.fore_color.rgb=rgb(color); sh.line.fill.background(); return sh
def pill(s,x,y,w,text,color=TEAL):
    box(s,x,y,w,.30,PANEL2)
    txt(s,x+.10,y+.035,w-.20,.23,text,10,color,True)
def node(s,x,y,w,h,title,sub='',color=TEAL):
    box(s,x,y,w,h,PANEL,line=color)
    txt(s,x+.16,y+.14,w-.32,.51,title,15.5,WHITE,True)
    if sub: txt(s,x+.16,y+.67,w-.32,h-.74,sub,12,MUTED)
def base(kicker,title,sub='',state='TARGET DESIGN'):
    s=prs.slides.add_slide(prs.slide_layouts[6]); s.background.fill.solid(); s.background.fill.fore_color.rgb=rgb(BG)
    txt(s,.48,.30,10,.22,kicker.upper(),10,TEAL,True)
    txt(s,.48,.79,12.35,.75,title,30,WHITE,True)
    if sub: txt(s,.50,1.64,12.2,.6,sub,16,MUTED)
    line(s,.5,7.02,12.83,7.02,LINE,.65)
    txt(s,.50,7.15,6,.16,'AMPHORA  /  INTERNAL DISCUSSION ONLY  /  16 SEP 2026',8,MUTED)
    txt(s,9.0,7.13,3.3,.20,state,8,TEAL,False,PP_ALIGN.RIGHT)
    txt(s,12.46,7.10,.36,.25,f'{len(prs.slides):02}',10,WHITE,True,PP_ALIGN.RIGHT)
    return s
def note(s,text,source='Graph_Native_Research_Engine.md; Product_Requirements.md'):
    s.notes_slide.notes_text_frame.text=(text+'\n\nINTERNAL ONLY. Architecture and illustrative workflows; not claims of validated product performance.\nSource: docs/Architecture/'+source)
def band(s,text,y=6.34,color=TEAL):
    box(s,.5,y,12.33,.43,PANEL2)
    txt(s,.68,y+.08,11.97,.3,text,14,color,True)

# 01 — editorial title with an abstract, editable research map.
s=base('Engine direction','Research pathways that adapt',state='INTERNAL CONCEPT BRIEF')
txt(s,.52,1.94,6.0,1.6,'One engine.\nMany scientific questions.',34,WHITE,True)
txt(s,.54,3.79,5.55,1.25,'Combine materials, chemistry and biology into a research map that can change as we learn.',21,MUTED)
pill(s,.55,5.61,2.1,'REUSABLE TOOLS'); pill(s,2.82,5.61,2.15,'ADAPTIVE PATHS',BLUE)
coords=[(7.08,2.05,'QUESTION',TEAL),(9.0,1.91,'KNOWLEDGE',PURPLE),(11.1,2.77,'ASSAY',BLUE),(8.35,3.73,'DESIGN',TEAL),(10.55,4.61,'EVIDENCE',GOLD),(7.13,5.57,'NEXT STEP',TEAL)]
for a,b in [(0,1),(0,3),(1,2),(1,3),(3,2),(3,4),(2,4),(4,5),(5,3)]:
    x,y,_,_=coords[a]; xx,yy,_,_=coords[b]; line(s,x+.58,y+.35,xx+.58,yy+.35,LINE,1.5)
for x,y,label,c in coords:
    box(s,x,y,1.42,.72,PANEL,c); txt(s,x+.10,y+.24,1.22,.28,label,11,c,True,PP_ALIGN.CENTER)
note(s,'The engine is the common research infrastructure. A campaign is the map of work for one question. Materials and biology are capability domains that can be composed; no single product dictates the engine. This deck explains intended behavior and the source-reviewed implementation baseline.')

# 02 — question to evidence.
s=base('The purpose','Turn a research question into an evidence trail','A campaign makes the next step, the reason for it, and the result visible.')
items=[('01','Define the question','What outcome matters?\nWhat must stay fixed?',TEAL),('02','Choose the tools','Select relevant data,\nmodels and experiments.',BLUE),('03','Learn and redirect','Test a candidate.\nRespond to the result.',PURPLE),('04','Explain the decision','What did we use?\nWhy this next step?',GOLD)]
for i,(n,t,b,c) in enumerate(items):
    x=.5+i*3.13; box(s,x,2.64,2.94,2.9,PANEL)
    txt(s,x+.20,2.84,2.5,.61,n,34,c,True)
    txt(s,x+.20,3.65,2.55,.77,t,20,WHITE,True)
    txt(s,x+.20,4.49,2.5,.8,b,16,MUTED)
    if i<3: arrow(s,x+2.95,4.0,x+3.11,4.0,c)
band(s,'The output includes the answer, the uncertainty, and the path that produced it.')
note(s,'Explain a campaign as a research map with executable steps. Each run retains evidence and decisions, including failures. The proposed engine supports human review and bounded automation; it does not treat model output as scientific qualification.')

# 03 — exact graph distinction.
s=base('Two maps and a record','Knowledge guides the work. Runs record what happened.','We author the scientific context and the research plan; execution creates the history.')
for x,t,tag,c in [(.5,'What we know','DOMAIN GRAPH',PURPLE),(4.75,'What we do','CAMPAIGN GRAPH',TEAL),(9,'What happened','DERIVED HISTORY',GOLD)]:
    box(s,x,2.65,3.83,3.22,PANEL); pill(s,x+.20,2.88,3.4,tag,c); txt(s,x+.20,3.40,3.43,.5,t,24,WHITE,True)
txt(s,.72,4.12,3.38,1.4,'Entities and mechanisms\nEvidence and contradictions\nLimits of our knowledge',18,MUTED)
txt(s,4.97,4.12,3.38,1.4,'Steps and decision points\nBranches and repeat loops\nHuman approvals',18,MUTED)
txt(s,9.22,4.12,3.38,1.4,'Inputs and results\nChanges and failures\nWho decided what',18,MUTED)
arrow(s,4.37,4.08,4.7,4.08,PURPLE); arrow(s,8.61,4.08,8.96,4.08,TEAL)
band(s,'New evidence can propose updates to either map; review determines what is accepted.')
note(s,'Domain and campaign graphs are separately authored and versioned. Execution/provenance is derived from runs, not a third independently authored truth. The UI should allow traversal in both directions between decisions and the knowledge they consumed.')

# 04 — dynamic flow with branch/retry/human gate.
s=base('Dynamic execution','The result determines where the campaign goes next','Illustrative delivery campaign: investigate, evaluate, branch, learn.')
node(s,.55,3.04,2.0,1.32,'Acquire data','Reference + project data',BLUE)
node(s,3.00,3.04,2.0,1.32,'Propose design','Inside fixed constraints')
node(s,5.45,3.04,2.0,1.32,'Evaluate','Models + measurements')
node(s,8.1,2.50,2.05,1.32,'Review','Evidence sufficient?',GOLD)
node(s,10.70,2.50,2.05,1.32,'Experiment','Approved next test',BLUE)
node(s,8.1,4.72,2.05,1.32,'Fill the gap','Acquire or measure',PURPLE)
for a,b in [(2.55,3),(5,5.45)]: arrow(s,a,3.70,b,3.70)
arrow(s,7.45,3.70,8.1,3.15); arrow(s,10.15,3.15,10.70,3.15,GOLD)
arrow(s,7.45,3.70,8.1,5.30,PURPLE)
txt(s,7.55,4.09,1.3,.3,'UNCERTAIN',10,PURPLE,True)
line(s,9.1,6.04,9.1,6.28,PURPLE); line(s,9.1,6.28,4,6.28,PURPLE); arrow(s,4,6.28,4,4.36,PURPLE)
line(s,11.72,2.50,11.72,2.24,BLUE); line(s,11.72,2.24,6.48,2.24,BLUE); arrow(s,6.48,2.24,6.48,3.04,BLUE)
txt(s,8.55,2.02,2.9,.25,'NEW MEASUREMENTS',10,BLUE,True)
txt(s,.57,5.15,5.75,.8,'Every loop has a purpose, a budget and a stopping rule.',23,WHITE,True)
note(s,'This is an illustrative target workflow. A graph determines composition and conditional behavior. It can request evidence, wait for human review or laboratory work, and resume. Changes require validated graph revisions; the diagram is not a claim that the graph interpreter is implemented.')

# 05 — mix and match tracks.
s=base('Cross-domain composition','Use the science the question needs','Polymers are one capability family. Biology can join the same campaign—or stand on its own.')
tracks=[('COATING','Material properties','Process design','Performance test',TEAL),('DRUG DELIVERY','Polymer + process','Release behavior','Biological response',BLUE),('BIOLOGY','Target + mechanism','Assay evidence','Next experiment',PURPLE)]
for i,(label,a,b,c,col) in enumerate(tracks):
    y=2.67+i*1.13; txt(s,.56,y+.29,2.00,.38,label,13,col,True)
    for j,t in enumerate([a,b,c]):
        x=2.75+j*3.35; box(s,x,y,2.93,.88,PANEL,line=col); txt(s,x+.16,y+.23,2.61,.47,t,17,WHITE,True)
        if j<2: arrow(s,x+2.95,y+.44,x+3.32,y+.44,col)
band(s,'Share capabilities when their inputs, scientific context and evidence are compatible.')
note(s,'Rows are examples of campaign composition, not fixed pipelines. A biology campaign must not require a polymer or epoxy node. A mixed campaign can connect materials to biological evidence only through typed, scientifically applicable adapters. Shared code does not automatically imply transferable calibration.')

# 06 — data scopes input fan-in.
s=base('Data and context','Bring the right information into each campaign','Source scope controls discovery and selection; permissions still apply to every use.')
for i,(title,desc,col) in enumerate([('DOMAIN','Shared scientific references',PURPLE),('PRODUCT','Target profile and requirements',BLUE),('PROJECT','Private experiments and measurements',TEAL)]):
    y=2.55+i*1.16; box(s,.53,y,3.62,.95,PANEL,line=col); txt(s,.73,y+.13,3.19,.27,title,12,col,True); txt(s,.73,y+.47,3.19,.47,desc,14,WHITE)
    arrow(s,4.2,y+.47,5.00,3.99,col)
node(s,5.03,3.05,2.55,1.90,'Acquire + prepare','Tables, papers, sequences,\nimages and instrument files',BLUE)
arrow(s,7.62,4,8.12,4)
node(s,8.16,3.05,2.10,1.90,'Retain','Versioned inputs\nand transformations',TEAL)
arrow(s,10.30,4,10.73,4)
node(s,10.77,3.05,2.02,1.90,'Review','Evidence ready\nfor the decision?',GOLD)
band(s,'A refreshed source creates a new version. Earlier decisions keep their original evidence trail.')
note(s,'Sources have reusable identities independent of campaigns. Acquisition and processing are campaign activities. Raw data is retained as permitted artifacts; extracted assertions pass review before trusted use. Snapshots, queries and transformations are pinned, and unavailable retention creates explicit replay limits. Credentials remain outside graph definitions.', 'Data_Sources_and_Acquisition.md; Product_Requirements.md')

# 07 — biological context field orbit.
s=base('Biological support','A result has meaning only in its biological context','The engine must keep observed activity, a predicted effect and an accepted conclusion distinct.')
box(s,.54,2.63,4.12,3.49,PANEL)
txt(s,.78,2.9,3.65,.86,'A biological observation',24,WHITE,True)
txt(s,.78,3.93,3.60,1.80,'What was measured?\nIn which system?\nUnder what conditions?\nWith what uncertainty?',20,BLUE)
fields=[('SYSTEM','Species, tissue, cell line'),('METHOD','Assay, controls, replicates'),('EXPOSURE','Dose, route, time'),('EVIDENCE','Source, limits, contradictions')]
for i,(h,b) in enumerate(fields):
    x=5.07+(i%2)*3.94; y=2.63+(i//2)*1.85
    box(s,x,y,3.76,1.64,PANEL); txt(s,x+.2,y+.18,3.34,.27,h,12,PURPLE,True); txt(s,x+.2,y+.70,3.32,.74,b,21,WHITE,True)
band(s,'Transferring a model across products, species or routes is a new applicability decision.',color=GOLD)
note(s,'Biological support includes entities, sequences, targets/pathways, systems, assays, samples and observations through versioned contracts. The UI must preserve uncertainty and contradictory evidence. Availability of a sequence adapter does not imply a validated biologics design system. No efficacy or safety claim is made here.', 'Product_Requirements.md#biological-and-cross-domain-research; Domain_Model.md')

# 08 — two evolution levels.
s=base('Learning at two levels','Improve the candidate—and eventually the research path','These are different search spaces with different validation needs.')
box(s,.52,2.56,5.95,3.38,PANEL); box(s,6.86,2.56,5.95,3.38,PANEL)
pill(s,.75,2.78,3.65,'PARAMETER SEARCH  /  KERNEL EXISTS'); pill(s,7.10,2.78,4.50,'PATHWAY EVOLUTION  /  FUTURE CAPABILITY',GOLD)
txt(s,.77,3.37,5.35,.51,'Find better settings',26,WHITE,True)
txt(s,7.10,3.37,5.35,.51,'Find a better sequence of work',26,WHITE,True)
for i,(label,p) in enumerate([('Composition',.66),('Process',.38),('Geometry',.78)]):
    y=4.2+i*.43; txt(s,.78,y,1.55,.29,label,13,MUTED); line(s,2.65,y+.13,5.90,y+.13,LINE,4); dot(s,2.65+p*3.25,y+.13,.17,TEAL)
for i,label in enumerate(['TEST','MODEL','REVIEW']):
    x=7.11+i*1.78; box(s,x,4.36,1.48,.60,PANEL2); txt(s,x+.08,4.54,1.32,.22,label,11,BLUE,True,PP_ALIGN.CENTER)
    if i<2: arrow(s,x+1.49,4.66,x+1.75,4.66,BLUE)
line(s,11.40,4.98,11.40,5.46,PURPLE); line(s,11.40,5.46,7.83,5.46,PURPLE); arrow(s,7.83,5.46,7.83,4.98,PURPLE)
band(s,'Evolution proposes alternatives. Constraints, evidence and human review decide what can run.')
note(s,'Existing NSGA-II performs numerical parameter evolution. The feature-selection GA is another implemented search space. Generic campaign-topology mutation and crossover were not found in the inspected paths. A reusable kernel still needs typed adapters, complete history and scientific evaluators; it is not production drug optimization.', 'Evolutionary_Code_Reuse_Baseline.md')

# 09 — human studio conceptual interface.
s=base('Human control','Scientists can see, question and redirect the work','Conceptual Graph Studio view: the map and the evidence remain connected.')
box(s,.53,2.50,12.28,3.72,PANEL,line=LINE)
box(s,.54,2.51,12.26,.42,PANEL2,r=0)
txt(s,.75,2.61,8,.2,'CAMPAIGN  /  DELIVERY STUDY  /  REVISION 04',10,MUTED,True)
pill(s,10.13,2.56,2.36,'WAITING FOR REVIEW',GOLD)
for x,y,w,h,t,b,c in [(1.0,3.55,2,1.23,'Reference data','Snapshot retained',BLUE),(3.61,3.55,2,1.23,'Evaluate design','Result available',TEAL),(6.22,3.55,2,1.23,'Review evidence','Decision pending',GOLD)]: node(s,x,y,w,h,t,b,c)
arrow(s,3.02,4.17,3.57,4.17); arrow(s,5.63,4.17,6.18,4.17)
txt(s,1.03,5.28,7.26,.5,'Inspect inputs     •     Compare revisions     •     Pause or revise',16,MUTED)
line(s,8.64,3.18,8.64,5.92,LINE,1)
txt(s,8.97,3.23,3.45,.40,'WHY THIS DECISION?',13,GOLD,True)
txt(s,8.97,3.88,3.45,1.55,'Evidence used\nAssumptions and gaps\nSupporting knowledge\nPrior attempts',18,WHITE)
band(s,'A human intervention becomes a recorded research event—not a lost conversation.')
note(s,'This is a conceptual UI, not an implemented product screenshot. Graph Studio is a core engine capability: editing, execution state, intervention, revision comparison, source inspection and traversal into the domain graph. User actions follow access, activation and evidence authority rules.')

# 10 — honest implementation baseline.
s=base('Starting point','We have reusable foundations. Scientific coverage must grow.','Source review distinguishes working algorithms from product-specific science.',state='SOURCE-REVIEWED BASELINE')
cols=[(.53,3.89,'REUSE',TEAL,['Numerical evolutionary search','Mixture and experiment grids','Evidence and replay patterns']), (4.72,3.89,'ADAPT',BLUE,['Product representations','Search-history integration','Feature-selection workflows']), (8.91,3.89,'BUILD + VALIDATE',GOLD,['Release / biological evaluators','Dynamic graph runtime + Studio','Campaign-pathway evolution'])]
for x,w,h,c,rows in cols:
    box(s,x,2.68,w,3.13,PANEL); txt(s,x+.2,2.95,w-.4,.35,h,17,c,True)
    for j,row in enumerate(rows): dot(s,x+.23,3.76+j*.70,.08,c); txt(s,x+.43,3.60+j*.70,w-.61,.60,row,16,WHITE)
band(s,'Reusable optimizer code does not establish a validated drug-release or biological model.',color=GOLD)
note(s,'The reviewed numerical kernel is amphora-solver-lab/src/nsga2.rs. EnginePool::mutate changes identity/lineage only; authorized chemistry generation separately exists but applies resin/hardener roles. Feature evolution is real but tied to NOMAD targets. No named-drug/PLGA/PK runtime evaluator was located in the searched source. Runtime tests were not rerun because toolchains/dependencies were unavailable. This is a static code baseline, not a readiness certification.', 'Evolutionary_Code_Reuse_Baseline.md')

# 11 — portfolio mapping, non-regulatory.
s=base('Five products, shared infrastructure','A portfolio of research maps—not five separate engines','The product brief supplies examples of different scientific work on a common foundation.')
headers=[(.73,2.47,3.70,'PRODUCT EXAMPLE'),(4.62,2.47,3.43,'POTENTIAL SHARED TOOLS'),(8.64,2.47,3.68,'EVIDENCE STAYS SPECIFIC')]
for x,y,w,t in headers: txt(s,x,y,w,.25,t,11,TEAL,True)
rows=[('Dexamethasone implant','Rod process + release analysis','Product, process and assay'),('Goserelin implant','Rod process + release analysis','Peptide stability + exposure'),('Once-daily buspirone','Mixture search + experiment design','Oral matrix + absorption'),('Veterinary analgesia depot','Depot analysis + evidence workflow','Route, species and duration'),('Canine OA injectable','Depot analysis + evidence workflow','Local/systemic branch + safety')]
for i,(a,b,c) in enumerate(rows):
    y=2.99+i*.56
    if i%2==0: box(s,.53,y-.07,12.27,.52,PANEL,r=0)
    txt(s,.73,y,3.68,.43,a,17,WHITE,True); txt(s,4.62,y,3.80,.43,b,15,MUTED); txt(s,8.64,y,3.91,.43,c,15,MUTED)
band(s,'Shared tools can reduce repeated engineering. Each product still has to earn its scientific evidence.')
note(s,'Product names are from AMPHORA_Top5_Product_Brief_1.docx. This slide maps software reuse, not market ranking, regulatory pathways, dates or clinical recommendations. Potential shared process/release modules need applicability checks. Cross-product and cross-species evidence does not transfer automatically. Buspirone tests a different delivery representation.', 'Evolutionary_Code_Reuse_Baseline.md; ../Artifacts/AMPHORA_Top5_Product_Brief_1.docx')

# 12 — learning loop with explicit gate.
s=base('Compounding learning','Keep the failures. Reuse the lessons. Review the transfer.','A completed campaign should improve the next research decision.')
positions=[(.62,2.82,'Design','Question + pathway',TEAL),(3.79,2.82,'Run','Results + failed attempts',BLUE),(6.96,2.82,'Compare','Evidence + alternatives',PURPLE),(10.13,2.82,'Review','Where does this apply?',GOLD)]
for x,y,t,b,c in positions: node(s,x,y,2.58,1.54,t,b,c)
for i in range(3): arrow(s,positions[i][0]+2.6,3.58,positions[i+1][0]-.05,3.58)
line(s,11.42,4.38,11.42,5.19,GOLD); arrow(s,11.42,5.19,7.30,5.19,GOLD)
box(s,3.96,4.82,3.28,.79,PANEL2,line=TEAL); txt(s,4.14,5.05,2.92,.35,'Reusable pathway',17,TEAL,True)
line(s,3.94,5.19,1.91,5.19,TEAL); arrow(s,1.91,5.19,1.91,4.38,TEAL)
band(s,'Retained history becomes useful only when context, rights and applicability travel with it.')
note(s,'The long-term aim is governed reuse of successful graph motifs and informative failures across projects. Raw customer/project evidence is private by default. Learned motifs remain candidate templates with ancestry, applicability and evidence; there is no automatic global promotion. This is a target capability.')

# 13 — phased deliverable roadmap.
s=base('Delivery sequence','Prove the shared engine in small, visible steps','Each gate demonstrates a new capability without claiming the next one is already complete.')
steps=[('1','Show portability','Rod + oral fixtures.\nOne search kernel.\nSynthetic objectives.',TEAL),('2','Make it executable','Versioned graphs.\nTyped tools + data.\nRetained state.',BLUE),('3','Make it scientific','Reviewed datasets.\nValidated models.\nClear assay context.',PURPLE),('4','Make it adaptive','Human Graph Studio.\nExperiment loops.\nThen evolve pathways.',GOLD)]
for i,(n,t,b,c) in enumerate(steps):
    x=.53+i*3.13; dot(s,x+.24,2.92,.31,c)
    if i<3: line(s,x+.43,2.92,x+3.16,2.92,LINE,2)
    txt(s,x,3.49,2.87,.36,'GATE '+n,12,c,True); txt(s,x,4.03,2.87,.86,t,23,WHITE,True); txt(s,x,5.02,2.86,1.08,b,16,MUTED)
band(s,'Biology acceptance includes a campaign that works without a polymer or epoxy node.')
note(s,'This is a logical delivery sequence, not a dated delivery promise. Minimal execution inspection should accompany the foundational fixture; richer Studio intervention follows. Scientific evaluators require held-out validation and appropriate applicability. The biological-support gates include materials-to-biology composition and a biology-only assay fixture.', 'Evolutionary_Code_Reuse_Baseline.md; Product_Requirements.md')

# 14 — concise internal alignment.
s=base('What we are aligning on','Build a research engine that can change its route','The scientific question determines the combination of tools.')
txt(s,.57,2.64,6.2,2.60,'Reusable science.\nConfigurable campaigns.\nVisible evidence.\nHuman control.',32,WHITE,True)
for i,(h,b,c) in enumerate([('REUSE THE FOUNDATION','Keep proven algorithms and evidence machinery.',TEAL),('EXPAND THE SCIENCE','Add biological and materials capabilities as needed.',BLUE),('PROVE THE CONNECTIONS','Validate the data, models and graph transitions.',GOLD)]):
    y=2.64+i*1.23; box(s,7.0,y,5.79,1.02,PANEL); txt(s,7.21,y+.14,5.35,.26,h,12,c,True); txt(s,7.21,y+.49,5.35,.44,b,16,WHITE)
band(s,'First proof: the same search core, two product representations, one inspectable evidence trail.')
note(s,'Internal discussion should select the first proof campaign and the evidence owners. The architecture supports polymers and biology without fixing a portfolio-wide pipeline. The source baseline is a guide to reuse and gaps, not an assertion that the complete engine or drug-product science is already delivered.', 'Evolutionary_Code_Reuse_Baseline.md; Graph_Native_Research_Engine.md')

prs.core_properties.title='Amphora | The Adaptive Research Engine'
prs.core_properties.subject='Internal non-technical engine architecture and reuse discussion'
prs.core_properties.author='Amphora'
prs.core_properties.keywords='INTERNAL ONLY; graph-native; biological research; materials; evolutionary search'
OUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(OUT)
print(f'Saved {len(prs.slides)} editable slides: {OUT}')
