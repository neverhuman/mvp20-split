"""Six investor-positioning slides plus a linked source appendix; editable shapes."""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN
import math

ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'docs/Artifacts/AMPHORA_Engine_VC_Positioning.pptx'
prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
BG='101D2C'; PANEL='192B3E'; PANEL2='22394D'; WHITE='F1F5F6'; MUTED='AFC1CE'
TEAL='56D9BD'; BLUE='83BFFF'; GOLD='F4C77E'; PURPLE='BEA5F4'; LINE='40586A'
SOURCES={
'S1':('Benchling AI','https://www.benchling.com/ai'),
'S2':('Citrine Catalyst','https://citrine.io/platform/citrine-catalyst/'),
'S3':('Citrine sequential learning','https://citrine.io/doe/'),
'S4':('Schrödinger active learning','https://www.schrodinger.com/platform/products/active-learning-applications/'),
'S5':('Synthace','https://www.synthace.com/what-we-do'),
'S6':('Automata LINQ','https://docs.automata.tech/tasks-transport'),
'S7':('Recursion OS','https://www.recursion.com/faq'),
'S8':('Claude Science','https://www.anthropic.com/news/claude-science-ai-workbench')}

def rgb(c): return RGBColor.from_string(c)
def box(s,x,y,w,h,fill=PANEL,stroke=None):
 sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h));sh.adjustments[0]=.12
 sh.fill.solid();sh.fill.fore_color.rgb=rgb(fill)
 if stroke: sh.line.color.rgb=rgb(stroke)
 else: sh.line.fill.background()
 return sh
def text(s,x,y,w,h,t,size=18,color=WHITE,bold=False,align=PP_ALIGN.LEFT,url=None):
 sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=sh.text_frame;tf.clear();tf.word_wrap=True
 tf.margin_top=tf.margin_bottom=tf.margin_left=tf.margin_right=0
 for i,row in enumerate(t.split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=row;p.font.name='Aptos';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=rgb(color);p.alignment=align;p.space_after=Pt(5)
  if url:
   for r in p.runs:r.hyperlink.address=url
 return sh
def line(s,x,y,xx,yy,c=LINE,width=1.5):
 sh=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x),Inches(y),Inches(xx),Inches(yy));sh.line.color.rgb=rgb(c);sh.line.width=Pt(width)
def arrow(s,x,y,xx,yy,c=TEAL):
 line(s,x,y,xx,yy,c,2);a=math.atan2(yy-y,xx-x)
 for d in [-.52,.52]:line(s,xx,yy,xx-.11*math.cos(a+d),yy-.11*math.sin(a+d),c,2)
def dot(s,x,y,c=TEAL,d=.12):
 sh=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x-d/2),Inches(y-d/2),Inches(d),Inches(d));sh.fill.solid();sh.fill.fore_color.rgb=rgb(c);sh.line.fill.background()
def chip(s,x,y,w,t,c=TEAL):
 box(s,x,y,w,.31,PANEL2);text(s,x+.10,y+.06,w-.2,.22,t,10,c,True)
def base(kicker,title,sub='',state='INVESTMENT THESIS'):
 s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=rgb(BG)
 text(s,.5,.30,12,.25,kicker.upper(),10,TEAL,True);text(s,.5,.80,12.3,.90,title,30,WHITE,True)
 if sub:text(s,.52,1.83,12.1,.56,sub,16,MUTED)
 line(s,.5,7.03,12.83,7.03,LINE,.65)
 text(s,.5,7.16,7,.19,'AMPHORA  /  CONFIDENTIAL INVESTOR DISCUSSION  /  SEPTEMBER 2026',8,MUTED)
 text(s,9.15,7.14,3.05,.22,state,8,TEAL,align=PP_ALIGN.RIGHT)
 text(s,12.45,7.10,.35,.29,f'{len(prs.slides):02}',10,WHITE,True,PP_ALIGN.RIGHT)
 return s
def band(s,t,c=TEAL,y=6.36):
 box(s,.52,y,12.29,.43,PANEL2);text(s,.72,y+.085,11.9,.31,t,14,c,True)
def notes(s,t,ids=()):
 body=t+'\n\nAmphora architecture is a proposed product direction. Existing numerical search/evidence code has been source-reviewed; full graph execution, biological evaluators and pathway evolution are not established as delivered. No commercial performance or exclusive novelty is claimed.\nLocal basis: docs/Architecture/Graph_Native_Research_Engine.md; Product_Requirements.md; Evolutionary_Code_Reuse_Baseline.md.\nCompetitive diligence: docs/Presentations/VC_Positioning_Research.md.\n'
 for key in ids:body+=f'\n[{key}] {SOURCES[key][0]}: {SOURCES[key][1]} (accessed 2026-09-16; vendor description)'
 s.notes_slide.notes_text_frame.text=body
def refs(s,ids,y=6.88):
 x=.54
 for key in ids:
  label=f'{key} {SOURCES[key][0]}'
  w=max(1.25,len(label)*.060)
  text(s,x,y,w,.14,label,7.5,MUTED,url=SOURCES[key][1]);x+=w+.17

# 1. Problem thesis: change is the thing the engine must represent.
s=base('The opportunity','Research changes. Its software should change with it.','Amphora is building an engine for research pathways that adapt as evidence arrives.')
text(s,.56,2.72,5.2,1.83,'Make the next step\nclear, executable\nand explainable.',33,WHITE,True)
text(s,.57,4.85,5.05,1.18,'Our thesis: less rebuilding between tools, stronger decision context, and more reusable work across programs.',20,MUTED)
chip(s,6.72,2.66,5.94,'ILLUSTRATIVE CAMPAIGN  /  PROPOSED ENGINE')
for x,y,t,c in [(6.8,3.39,'DESIGN',TEAL),(9.15,3.39,'TEST',BLUE),(11.22,3.39,'REVIEW',GOLD),(8.97,5.04,'NEW EVIDENCE',PURPLE)]:
 w=1.49 if t!='NEW EVIDENCE' else 2.72
 box(s,x,y,w,.81,PANEL,c);text(s,x+.10,y+.28,w-.2,.26,t,11,c,True,PP_ALIGN.CENTER)
arrow(s,8.31,3.79,9.09,3.79);arrow(s,10.66,3.79,11.16,3.79)
line(s,11.97,4.22,11.97,5.44,GOLD);arrow(s,11.97,5.44,11.71,5.44,GOLD)
line(s,8.94,5.44,7.55,5.44,PURPLE);arrow(s,7.55,5.44,7.55,4.24,PURPLE)
band(s,'The economic target: reduce the work between a promising result and a defensible next decision.')
notes(s,'Open on a buyer problem to validate, not a fabricated market statistic. Research objectives, evidence and constraints change. The proposed engine represents that changing decision process as an executable research map. The value hypotheses must be measured against the current workflow and credible incumbent/custom-integration alternatives.')

# 2. Specific differentiation architecture, not generic AI.
s=base('The differentiation thesis','Make the research pathway an asset we can improve','One connected system for scientific knowledge, executable work and outcome-linked learning.',state='TARGET ARCHITECTURE')
cards=[(.54,'01','KNOWLEDGE','What can we rely on?','Mechanisms, evidence,\nuncertainty and limits.',PURPLE),(4.78,'02','CAMPAIGN','What should run next?','Editable steps, branches,\nexperiments and human gates.',TEAL),(9.02,'03','HISTORY','What did we learn?','Every attempt, decision\nand outcome retained.',GOLD)]
for x,n,k,h,b,c in cards:
 box(s,x,2.73,3.78,2.70,PANEL);text(s,x+.18,2.94,3.4,.30,n+'  '+k,13,c,True)
 text(s,x+.18,3.58,3.40,.76,h,24,WHITE,True);text(s,x+.18,4.54,3.38,.67,b,17,MUTED)
arrow(s,4.37,4.05,4.72,4.05,PURPLE);arrow(s,8.61,4.05,8.96,4.05,TEAL)
text(s,.74,5.71,11.93,.38,'FUTURE SEARCH: compare and evolve whole research pathways under scientific and human review.',16,TEAL,True)
band(s,'Our proposed distinction is this governed combination across materials, chemistry and biology.')
notes(s,'There are two authored graphs: domain knowledge and executable campaigns. History/provenance is derived from runs. Target novelty is the integrated representation plus governed evolution of campaign topology, not merely a canvas, a chatbot or candidate optimization. This is a differentiation hypothesis, not evidence that no competitor has similar features or that the architecture is patentable. Existing code supports numerical/feature evolution; general pathway evolution remains future work.')

# 3. Fair named-competitor overview with no unsupported checkmark claims.
s=base('Competitive landscape','Strong alternatives raise the bar for Amphora','These products already solve meaningful parts of scientific work. We must win on the combined use case.',state='VENDOR-SOURCED COMPARISON')
text(s,.69,2.60,3.4,.25,'EXISTING ALTERNATIVE',11,TEAL,True);text(s,4.04,2.60,4.73,.25,'WHAT IT ALREADY BRINGS',11,TEAL,True)
rows=[('Benchling / Claude Science','Scientific AI, context and tool orchestration','S1','S8'),('Citrine / Schrödinger','Materials or molecular optimization','S2','S4'),('Synthace / Automata','Experiment design, automation and task graphs','S5','S6'),('Recursion OS','Integrated computational + physical discovery','S7',None)]
for i,(a,b,k,k2) in enumerate(rows):
 y=3.12+i*.65;box(s,.54,y-.1,8.29,.58,PANEL)
 text(s,.71,y,3.15,.46,a,16,WHITE,True,url=SOURCES[k][1]);text(s,4.04,y,4.51,.46,b,15,MUTED)
box(s,9.17,2.72,3.64,3.20,PANEL2,TEAL)
text(s,9.38,2.95,3.20,.32,'AMPHORA TARGET',12,TEAL,True)
text(s,9.38,3.55,3.18,1.26,'Versioned, evolvable\nresearch maps linked\nto scientific evidence.',22,WHITE,True)
text(s,9.38,5.04,3.18,.72,'Cross-domain.\nScientist-controlled.',17,MUTED)
band(s,'The test: better campaign change, scientific traceability and reuse—not a longer feature list.')
refs(s,['S1','S2','S4','S5','S6','S7','S8'])
notes(s,'These rows are summaries of advertised strengths, not mutually exclusive categories or claims about missing features. Citrine Catalyst already creates AI workflows and supports reviewed domain-knowledge edits. Automata explicitly documents graph tasks and conditionals. Benchling and Claude Science advertise scientific assistance and traceability. Recursion describes end-to-end discovery. Direct evaluation is needed to establish whether Amphora offers a differentiated bundle. We must not say only Amphora uses graphs, connects wet/dry research, retains lineage, or supports human review.', ['S1','S2','S3','S4','S5','S6','S7','S8'])

# 4. Why it must be engineered: capability changing science and evidence.
s=base('Why this engine needs to be built','A change of scientific direction should preserve the work','Our design makes the pathway, its scientific limits and its history part of the same product.',state='DESIGN GOALS / TO PROVE')
for i,(n,h,b,c) in enumerate([
 ('01','Change the path','Insert a new assay or evaluator\nwithout rebuilding the whole campaign.',TEAL),
 ('02','Protect the meaning','Carry source, species, method and\nmodel limits into every decision.',BLUE),
 ('03','Reuse what worked','Transfer reviewed steps and patterns\nwhile keeping product-specific evidence.',PURPLE)]):
 y=2.74+i*1.03;box(s,.55,y,7.52,.90,PANEL)
 text(s,.75,y+.24,.58,.42,n,21,c,True);text(s,1.52,y+.10,2.50,.74,h,19,WHITE,True);text(s,4.17,y+.16,3.68,.64,b,14,MUTED)
box(s,8.52,2.74,4.28,2.97,PANEL2)
text(s,8.76,3.02,3.8,.31,'WHY NOW',13,GOLD,True)
text(s,8.76,3.61,3.75,1.95,'Scientific AI and lab automation are expanding. Adaptable research execution can turn those tools into repeatable value.',19,WHITE)
band(s,'Build the shared execution and evidence layer; connect specialist models, data platforms and labs.')
refs(s,['S1','S2','S6','S8'])
notes(s,'Why now is an inference from current vendor offerings in scientific AI, orchestration and automation, not a measured market-size trend. The proposed engineering work is a versioned interpreter, scientific applicability contracts, retained state/history, intervention controls and evaluation integration. A generic workflow runner can schedule tasks; assembling the required scientific behavior still requires implementation. Existing alternatives may satisfy some or all buyer needs; the build decision must be justified by a measured pilot rather than categorical dismissal.', ['S1','S2','S6','S8'])

# 5. Moat, rights and focused entry.
s=base('Potential defensibility','Each campaign can leave behind more than a result','The accumulating asset is a reviewed history of which research paths worked, where and why.',state='MOAT + COMMERCIAL HYPOTHESIS')
stages=[(.56,'RUN','Attempts + outcomes',TEAL),(3.77,'COMPARE','Cost + evidence + context',BLUE),(6.98,'REVIEW','Applicability + rights',PURPLE),(10.19,'REUSE','Campaign patterns',GOLD)]
for x,h,b,c in stages:
 box(s,x,2.78,2.59,1.45,PANEL,c);text(s,x+.17,2.99,2.25,.31,h,13,c,True);text(s,x+.17,3.55,2.25,.58,b,18,WHITE,True)
for i in range(3):arrow(s,stages[i][0]+2.62,3.52,stages[i+1][0]-.06,3.52)
line(s,11.48,4.27,11.48,4.61,GOLD);line(s,11.48,4.61,1.85,4.61,GOLD);arrow(s,1.85,4.61,1.85,4.27,GOLD)
text(s,.59,5.08,5.9,.44,'START WITH A FOCUSED PRODUCT PROGRAM',12,TEAL,True)
text(s,.59,5.57,5.88,.54,'Prove value on delivery/formulation campaigns.',19,WHITE,True)
text(s,7.09,5.08,5.68,.44,'EXPAND THROUGH REPEATABLE CAPABILITIES',12,BLUE,True)
text(s,7.09,5.57,5.68,.54,'Reuse across materials and biology.',19,WHITE,True)
band(s,'Defensibility must be earned through validated outcomes, integration and authorized reuse.')
notes(s,'A potential moat is outcome-linked campaign history plus validated capability adapters and integration into repeat workflows. The architecture alone is not a moat. Private project evidence cannot be pooled without rights; no universal network effect is assumed. Proposed commercial sequence: focused sponsored/product development programs, then validate repeatable platform or licensing demand. This is a hypothesis, not current customer, revenue or business-model evidence. The five-product brief informs examples without establishing their regulatory or market claims.')

# 6. Capital-to-proof: factual code baseline and milestones, no fake round.
s=base('What investment would prove','Convert reusable code into a validated research engine','A staged proof program connects the architecture to customer value.',state='PROPOSED MILESTONES')
box(s,.55,2.72,3.56,3.29,PANEL)
text(s,.77,2.99,3.11,.28,'FOUNDATION IN THE CHECKOUT',12,TEAL,True)
text(s,.77,3.58,3.10,1.77,'Evolutionary search\nMixture optimization\nFeature selection\nEvidence machinery',22,WHITE,True)
text(s,.77,5.38,3.1,.42,'Source-reviewed; integration remains.',13,MUTED)
milestones=[('A','Working engine','Two product representations; one\nsearch core and traceable graph.'),('B','Scientific credibility','Reviewed data and validated evaluators;\nexplicit limits on transfer.'),('C','Commercial evidence','Pilot against the current workflow;\nmeasure change effort, reuse and cost.')]
for i,(n,h,b) in enumerate(milestones):
 y=2.72+i*1.11;box(s,4.55,y,8.25,.96,PANEL)
 text(s,4.77,y+.25,.45,.36,n,22,[TEAL,BLUE,GOLD][i],True)
 text(s,5.51,y+.13,2.47,.70,h,21,WHITE,True)
 text(s,8.31,y+.20,4.23,.60,b,15,MUTED)
band(s,'Investment case: prove repeatable research value first; earn the right to expand the platform.')
notes(s,'No funding amount, revenue, partnerships or advantage percentage is supplied. Proposed first proof uses rod-depot and oral-matrix synthetic fixtures for software portability, followed separately by scientific validation and a cross-domain biological fixture. Generic graph runtime, full Studio, drug-specific science and pathway evolution remain work to do. Measure setup/change time, decision reconstruction, reuse, gate behavior and cost per reviewed decision under pre-agreed baselines. Do not describe historical source checks as current runtime certification.')

# 7. Appendix: sources are clickable, with scope visible.
s=base('Appendix / evidence base','Competitive sources and claim boundaries','Official product descriptions reviewed 16 September 2026. Links open the underlying source.',state='DILIGENCE APPENDIX')
labels={'S1':'Scientific data, AI and model access','S2':'Goal-to-workflow AI in materials','S3':'Iterative experimental optimization','S4':'Physics-based active learning','S5':'Experiment design and automation','S6':'Task graphs and conditional execution','S7':'Integrated drug-discovery platform','S8':'Scientific agents and reproducible work'}
for i,(key,(name,url)) in enumerate(SOURCES.items()):
 x=.57+(i//4)*6.32;y=2.67+(i%4)*.77
 text(s,x,y,5.88,.32,f'{key}  {name}',17,TEAL,True,url=url)
 text(s,x,y+.37,5.88,.30,labels[key],13,MUTED)
box(s,.54,6.00,12.26,.79,PANEL2)
text(s,.73,6.16,11.89,.55,'Vendor claims are not independent tests. Amphora differentiation is a thesis to validate;\nthis deck makes no exclusive-novelty, patentability, clinical-success or measured-ROI claim.',13,WHITE)
notes(s,'Use this appendix for diligence. Competitor rows are not exhaustive comparisons and do not imply that omitted capabilities are absent. All evidence is from official sources; our intended architecture is from the local specifications. Local companion memo VC_Positioning_Research.md records the claims, boundaries and pilot questions.',tuple(SOURCES))

prs.core_properties.title='Amphora | Investor Engine Positioning'
prs.core_properties.subject='Research engine differentiation, alternatives and investment milestones'
prs.core_properties.author='Amphora'
prs.core_properties.keywords='Confidential; investor discussion; research engine; competitive positioning'
prs.save(OUT)
# Match clickable source links to the dark presentation theme.
from zipfile import ZipFile, ZIP_DEFLATED
from io import BytesIO
from lxml import etree
with ZipFile(OUT) as zin:
 payloads={name:zin.read(name) for name in zin.namelist()}
for name,data in list(payloads.items()):
 if name.startswith('ppt/theme/') and name.endswith('.xml'):
  root=etree.fromstring(data)
  for key in ['hlink','folHlink']:
   for node in root.findall('.//{http://schemas.openxmlformats.org/drawingml/2006/main}'+key):
    for child in node: child.set('val',TEAL)
  payloads[name]=etree.tostring(root,xml_declaration=True,encoding='UTF-8',standalone=True)
with ZipFile(OUT,'w',ZIP_DEFLATED) as zout:
 for name,data in payloads.items():zout.writestr(name,data)
print(f'Saved {len(prs.slides)} slides: {OUT}')
